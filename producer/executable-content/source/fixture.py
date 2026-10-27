"""Catalog-authored offline fixtures. No package acquisition or paid/backend invocation."""
import sys
import json
import importlib.metadata
import signal
import platform
import re
from pathlib import Path


def main():
    def timeout_handler(*_args):
        raise TimeoutError('fixture-timeout')

    signal.signal(signal.SIGTERM, timeout_handler)
    if hasattr(signal, 'SIGALRM'):
        signal.signal(signal.SIGALRM, timeout_handler)
        signal.alarm(25)
    workflow, destination = sys.argv[1:]
    requirements = json.loads(Path(__file__).with_name('requirements.json').read_text())
    selected = next(row for row in requirements['workflows'] if row['id'] == workflow)
    actual = '.'.join(map(str, sys.version_info[:3]))
    if actual != selected['python']:
        raise ValueError('runtime-version-mismatch: requires Python ' + selected['python'])
    if sys.prefix == sys.base_prefix:
        raise ValueError('isolated-environment-required')
    if sys.platform != 'linux' or platform.machine().lower() not in ('x86_64', 'amd64'):
        raise ValueError('platform-lock-unavailable: only linux-x64 resolution is retained')
    lock = Path(__file__).with_name('locks') / (workflow + '-linux-x64.txt')
    if not lock.is_file():
        raise ValueError('dependency-closure-unavailable')
    for package, version in re.findall(r'^([A-Za-z0-9_.-]+)==([^\s]+)', lock.read_text(), re.MULTILINE):
        if importlib.metadata.version(package) != version:
            raise ValueError('dependency-version-mismatch: ' + package + '==' + version)
    for package, version in selected['packages'].items():
        if importlib.metadata.version(package) != version:
            raise ValueError('dependency-version-mismatch: ' + package + '==' + version)
    output = Path(destination)
    if output.exists() or output.is_symlink():
        raise ValueError('existing-output-refused')
    # Exclusive creation protects adjacent user files; each fixture has its own new directory.
    output.mkdir()
    if workflow == 'mcp-evaluator':
        from mcp.types import CallToolResult, TextContent
        result = CallToolResult(content=[TextContent(type='text', text='fixture answer')])
        (output / 'evaluation.json').write_text(result.model_dump_json())
    elif workflow == 'skill-authoring':
        import yaml
        document = yaml.safe_load('name: fixture-skill\ndescription: Offline validation fixture\n')
        if document['name'] != 'fixture-skill':
            raise ValueError('invalid-skill')
        (output / 'validation.json').write_text(json.dumps(document))
    elif workflow == 'evalview':
        # No model call: package/import validation only, never claimed as evaluation success.
        import evalview
        (output / 'validation.json').write_text(json.dumps({'imported': evalview.__name__, 'modelInvoked': False}))
    elif workflow == 'gif':
        from PIL import Image
        frames = [Image.new('RGB', (32, 32), color) for color in ['red', 'blue']]
        frames[0].save(output / 'fixture.gif', save_all=True, append_images=frames[1:], duration=100, loop=0)
        with Image.open(output / 'fixture.gif') as image:
            if image.n_frames != 2 or image.size != (32, 32):
                raise ValueError('invalid-gif')
    elif workflow == 'slides':
        from reportlab.pdfgen import canvas
        from pptx import Presentation
        slides = Presentation()
        slides.slides.add_slide(slides.slide_layouts[6])
        slides.save(output / 'fixture.pptx')
        if len(Presentation(output / 'fixture.pptx').slides) != 1:
            raise ValueError('invalid-pptx')
        pdf = canvas.Canvas(str(output / 'fixture.pdf'), invariant=1)
        pdf.drawString(72, 720, 'Offline document fixture')
        pdf.save()
        if not (output / 'fixture.pdf').read_bytes().startswith(b'%PDF-'):
            raise ValueError('invalid-pdf')
    elif workflow == 'browser':
        browser_fixture(output)
    print(json.dumps({'status': 'fixture-complete', 'workflow': workflow, 'nativeAdmission': 'unverified'}))


def browser_fixture(output):
    # Loopback only; no shell, arbitrary server argv, fixed port or terminal launch.
    import threading
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    from playwright.sync_api import sync_playwright

    if sys.platform == 'win32':
        raise ValueError('native-windows-browser-process-cleanup-unqualified')

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200)
            self.send_header('Content-Type', 'text/html; charset=utf-8')
            self.end_headers()
            self.wfile.write(b'<title>Catalog fixture</title><h1>Local page</h1>')

        def log_message(self, *_args):
            pass

    server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        with sync_playwright() as playwright:
            # Only the explicitly acquired Playwright-managed browser, no Chrome/PATH fallback.
            browser = playwright.chromium.launch(headless=True, chromium_sandbox=True)
            try:
                page = browser.new_page()
                page.route('**/*', lambda route: route.continue_() if route.request.url.startswith(
                    'http://127.0.0.1:' + str(server.server_port) + '/') else route.abort())
                page.goto('http://127.0.0.1:' + str(server.server_port), timeout=5000)
                if page.title() != 'Catalog fixture':
                    raise ValueError('local-page-failed')
                page.screenshot(path=str(output / 'local-page.png'))
                (output / 'browser.json').write_text(json.dumps({'version': browser.version, 'route': 'playwright-managed-chromium'}))
            finally:
                browser.close()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)
        if thread.is_alive():
            raise ValueError('server-cleanup-failed')


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print(json.dumps({'status': 'blocked', 'reason': str(error)}))
        sys.exit(2)
